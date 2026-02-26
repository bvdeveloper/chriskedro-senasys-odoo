#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate Release Note as Word document (.docx). Run from repo root or docs/."""

import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_para(doc, text, bold=False, style=None):
    p = doc.add_paragraph(text, style=style)
    if bold:
        for run in p.runs:
            run.bold = True
    return p

def main():
    script_dir = os.path.dirname(os.path.abspath(__file__))
    out_path = os.path.join(script_dir, "Release_Note_Drawing_Revision_RFQ_Attachments.docx")

    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    # Title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('RELEASE NOTE')
    r.bold = True
    r.font.size = Pt(16)
    doc.add_paragraph()

    add_para(doc, 'Product: Senasys Odoo Customization (chriskedro-senasys-odoo)', bold=True)
    add_para(doc, 'Module: bv_senasys_custom_fields', bold=True)
    add_para(doc, 'Release Date: February 2025', bold=True)
    add_para(doc, 'Version: 15.0', bold=True)
    doc.add_paragraph()

    add_para(doc, 'Subject: Drawing Revision Field, Product PDF & Revision Note Attachments on Purchase RFQ/PO Email', bold=True)
    doc.add_paragraph()
    doc.add_paragraph('_' * 80)
    doc.add_paragraph()

    # 1. OVERVIEW
    add_para(doc, '1. OVERVIEW', bold=True)
    doc.add_paragraph(
        'This release introduces enhancements to the Product and Purchase Order modules to support '
        'drawing revision tracking and automatic attachment of product drawings and revision notes '
        'when sending Request for Quotation (RFQ) or Purchase Order (PO) emails to vendors.')
    doc.add_paragraph()

    # 2. FEATURES
    add_para(doc, '2. FEATURES IMPLEMENTED', bold=True)
    doc.add_paragraph()

    add_para(doc, '2.1 Drawing Revision Field on Product Form', bold=True)
    doc.add_paragraph('• A new field "Drawing Revision" has been added to the Product form.')
    doc.add_paragraph('• Location: Product → Technical tab → Vendor Drawings/Docs section, placed immediately before the existing "PDF 1_1" field.')
    doc.add_paragraph('• Field type: Text (multi-line).')
    doc.add_paragraph('• Purpose: Allows users to store revision notes or revision identifiers for product drawings. This value is used to generate the revision note text file attached to RFQ/PO emails.')
    doc.add_paragraph()

    add_para(doc, '2.2 Automatic Attachment of Product PDF (pdf_1_1) in RFQ/PO Email', bold=True)
    doc.add_paragraph('• When you send an RFQ or a Purchase Order by email ("Send by Email" or "Send PO by Email"), the system now automatically attaches the "PDF 1_1" file from each product present on the purchase order lines.')
    doc.add_paragraph('• If multiple order lines use the same product, the PDF is attached once per product (no duplicates).')
    doc.add_paragraph('• Only products that have a "PDF 1_1" file uploaded will have their PDF attached.')
    doc.add_paragraph('• The standard PO/RFQ report (e.g. PO_P01789.pdf) continues to be attached as before; product PDFs are added in addition to it.')
    doc.add_paragraph()

    add_para(doc, '2.3 Dynamic Revision Note (TXT) File per Product', bold=True)
    doc.add_paragraph('• For each distinct product on the purchase order lines, the system generates a text file containing the "Drawing Revision" value of that product.')
    doc.add_paragraph('• File naming: revision_note_<Product Part/Name>.txt (e.g. revision_note_ABC-123.txt).')
    doc.add_paragraph('• Content: The exact text stored in the product\'s "Drawing Revision" field. If the field is empty, no revision note file is created for that product (only the PDF is attached if available).')
    doc.add_paragraph('• These revision note files are attached to the same RFQ/PO email along with the product PDFs and the main PO/RFQ report.')
    doc.add_paragraph()

    # 3. TECHNICAL SUMMARY
    add_para(doc, '3. TECHNICAL SUMMARY', bold=True)
    doc.add_paragraph('• New field: drawing_revision (Text) on product.template and product.product.')
    doc.add_paragraph('• Product form view updated to display "Drawing Revision" before "PDF 1_1" in the Vendor Drawings/Docs section.')
    doc.add_paragraph('• Purchase Order: action_rfq_send overridden to pass product-based attachments (revision TXT + pdf_1_1) into the email composer context for both RFQ and PO emails.')
    doc.add_paragraph('• New helper: _get_rfq_product_attachments on purchase.order builds the list of temporary attachments (revision note TXT and pdf_1_1) per distinct product on the order lines.')
    doc.add_paragraph('• Mail Composer: mail.compose.message extended so that when composing an email for a Purchase Order, the product attachments are merged with the template attachments, ensuring they appear in the Compose Email dialog and are sent with the email.')
    doc.add_paragraph()

    # 4. USER INSTRUCTIONS
    add_para(doc, '4. USER INSTRUCTIONS', bold=True)
    doc.add_paragraph()

    add_para(doc, '4.1 Setting Up Products', bold=True)
    doc.add_paragraph('1. Open the Product form (Product → select or create a product).')
    doc.add_paragraph('2. Go to the "Technical" tab.')
    doc.add_paragraph('3. In the "Vendor Drawings/Docs" section:')
    doc.add_paragraph('   • Enter the revision information in the new "Drawing Revision" field (e.g. revision number or notes).')
    doc.add_paragraph('   • Upload the drawing file in "PDF 1_1" and set "PDF 1_1 filename" if desired.')
    doc.add_paragraph('4. Save the product.')
    doc.add_paragraph()

    add_para(doc, '4.2 Sending RFQ or PO by Email', bold=True)
    doc.add_paragraph('1. Create or open a Purchase Order (RFQ or confirmed PO) and add order lines with the products that have Drawing Revision and/or PDF 1_1 set.')
    doc.add_paragraph('2. Click "Send by Email" (for RFQ) or "Send PO by Email" (for confirmed PO).')
    doc.add_paragraph('3. The Compose Email dialog will open with attachments pre-filled:')
    doc.add_paragraph('   • The standard PO/RFQ report (e.g. PO_P01789.pdf).')
    doc.add_paragraph('   • For each product: its revision note file (revision_note_<product>.txt), if Drawing Revision is filled.')
    doc.add_paragraph('   • For each product: its PDF 1_1 file, if uploaded.')
    doc.add_paragraph('4. Review the recipients and body, then click "Send" to send the email with all attachments.')
    doc.add_paragraph()

    # 5. NOTES
    add_para(doc, '5. NOTES FOR CLIENT', bold=True)
    doc.add_paragraph('• Attachments created for the email (revision TXT and product PDFs) are temporary and are cleaned up by Odoo\'s mail composer after use; they do not clutter the database.')
    doc.add_paragraph('• If the same product appears on multiple lines of the same order, only one revision note and one PDF are attached for that product.')
    doc.add_paragraph('• This behavior applies to both "Send by Email" (RFQ) and "Send PO by Email" (confirmed Purchase Order).')
    doc.add_paragraph()

    doc.add_paragraph('_' * 80)
    doc.add_paragraph()
    add_para(doc, 'Document prepared by: Brainvire', bold=True)
    add_para(doc, 'For: Senasys – Drawing Revision & RFQ/PO Email Attachments Feature', bold=True)

    doc.save(out_path)
    print('Created:', out_path)
    return out_path

if __name__ == '__main__':
    main()
